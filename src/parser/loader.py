"""Document loader supporting PDF, DOCX, and TXT formats with spatial layout ordering."""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Union, BinaryIO, Optional, Dict, Any, List
import io
import logging

logger = logging.getLogger(__name__)


@dataclass
class LoadedDocument:
    raw_text: str
    file_type: str
    page_count: int = 1
    pages: List[str] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)
    filename: Optional[str] = None
    is_scanned: bool = False


class DocumentLoader:
    """Universal loader for PDF, DOCX, and TXT resume files with layout-aware spatial sorting."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

    @classmethod
    def load_file(cls, file_path: Union[str, Path]) -> LoadedDocument:
        """Load document from a file path."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        suffix = path.suffix.lower()
        if suffix not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file format '{suffix}'. Supported: {', '.join(cls.SUPPORTED_EXTENSIONS)}")

        if suffix == ".pdf":
            return cls._load_pdf(path)
        elif suffix == ".docx":
            return cls._load_docx(path)
        elif suffix == ".txt":
            return cls._load_txt(path)
        else:
            raise ValueError(f"Unsupported file format '{suffix}'")

    @classmethod
    def load_bytes(cls, data: bytes, filename: str) -> LoadedDocument:
        """Load document from in-memory byte buffer."""
        suffix = Path(filename).suffix.lower()
        if suffix not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file format '{suffix}'. Supported: {', '.join(cls.SUPPORTED_EXTENSIONS)}")

        stream = io.BytesIO(data)
        if suffix == ".pdf":
            return cls._load_pdf_stream(stream, filename)
        elif suffix == ".docx":
            return cls._load_docx_stream(stream, filename)
        elif suffix == ".txt":
            return cls._load_txt_bytes(data, filename)
        else:
            raise ValueError(f"Unsupported file format '{suffix}'")

    # --- PDF Processing with Spatial Column Awareness ---

    @staticmethod
    def _extract_page_blocks_layout_aware(page) -> str:
        """Extract text from PyMuPDF page preserving 2-column spatial reading order.
        
        Standard page.get_text("text") reads in raw stream order, which often mashes
        left-hand sidebars and right-hand body paragraphs together line-by-line.
        This method analyzes block bounding boxes:
        - If two distinct column clusters are detected (e.g. sidebar on left, body on right),
          it sorts left-column blocks top-down, then right-column blocks top-down.
        - Otherwise, it uses PyMuPDF's spatial sort=True reading order.
        """
        try:
            blocks = page.get_text("blocks", sort=True)
            if not blocks:
                return page.get_text("text") or ""

            # Filter text blocks (block_type == 0) that have non-empty text
            text_blocks = [b for b in blocks if len(b) >= 5 and b[4].strip() and (len(b) < 7 or b[6] == 0)]
            if not text_blocks:
                return page.get_text("text") or ""

            page_width = page.rect.width
            mid_thresh = page_width * 0.48

            left_blocks = []
            right_blocks = []
            span_blocks = []

            for b in text_blocks:
                x0, y0, x1, y1, text = b[0], b[1], b[2], b[3], b[4]
                # Full width span (e.g. candidate name header across top or footer)
                if x0 < page_width * 0.25 and x1 > page_width * 0.65:
                    span_blocks.append(b)
                elif x1 <= mid_thresh:
                    left_blocks.append(b)
                elif x0 >= page_width * 0.35:
                    right_blocks.append(b)
                else:
                    span_blocks.append(b)

            # If both columns have significant blocks (e.g. 2+ each), preserve column integrity
            if len(left_blocks) >= 2 and len(right_blocks) >= 2:
                col_top = min(min(b[1] for b in left_blocks), min(b[1] for b in right_blocks))
                col_bottom = max(max(b[3] for b in left_blocks), max(b[3] for b in right_blocks))

                top_spans = [b for b in span_blocks if b[3] <= col_top + 20]
                bottom_spans = [b for b in span_blocks if b[1] >= col_bottom - 20]
                mid_spans = [b for b in span_blocks if b not in top_spans and b not in bottom_spans]

                top_spans.sort(key=lambda b: (b[1], b[0]))
                left_blocks.sort(key=lambda b: (b[1], b[0]))
                right_blocks.sort(key=lambda b: (b[1], b[0]))
                bottom_spans.sort(key=lambda b: (b[1], b[0]))

                ordered_blocks = top_spans + mid_spans + left_blocks + right_blocks + bottom_spans
                return "\n\n".join(b[4].strip() for b in ordered_blocks if b[4].strip())

            # Default to spatially sorted blocks
            return "\n\n".join(b[4].strip() for b in text_blocks if b[4].strip())
        except Exception as e:
            logger.debug(f"Layout-aware block extraction failed, falling back to get_text: {e}")
            return page.get_text("text") or ""

    @classmethod
    def _load_pdf(cls, path: Path) -> LoadedDocument:
        """Extract text from PDF using PyMuPDF (fitz) with layout-awareness and pdfplumber fallback."""
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(str(path))
            pages_text = []
            for page in doc:
                text = cls._extract_page_blocks_layout_aware(page)
                pages_text.append(text)
            page_count = len(doc)
            doc.close()
            full_text = "\n\n".join(pages_text)
            
            # Check for scanned PDF
            total_chars = sum(len(p.strip()) for p in pages_text)
            is_scanned = (total_chars < 50 and page_count >= 1)

            metadata = {"engine": "pymupdf_spatial"}
            if is_scanned:
                metadata["is_scanned"] = True
                metadata["scanned_warning"] = "Low text density: document may be a scanned image or photograph."

            return LoadedDocument(
                raw_text=full_text,
                file_type="pdf",
                page_count=page_count,
                pages=pages_text,
                metadata=metadata,
                filename=path.name,
                is_scanned=is_scanned
            )
        except Exception as e:
            logger.warning(f"PyMuPDF failed on {path.name}: {e}. Trying pdfplumber fallback.")

        # Fallback to pdfplumber
        try:
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                pages_text = [page.extract_text() or "" for page in pdf.pages]
                full_text = "\n\n".join(pages_text)
                total_chars = sum(len(p.strip()) for p in pages_text)
                is_scanned = (total_chars < 50 and len(pdf.pages) >= 1)
                
                metadata = {"engine": "pdfplumber"}
                if is_scanned:
                    metadata["is_scanned"] = True
                    metadata["scanned_warning"] = "Low text density: document may be a scanned image."

                return LoadedDocument(
                    raw_text=full_text,
                    file_type="pdf",
                    page_count=len(pdf.pages),
                    pages=pages_text,
                    metadata=metadata,
                    filename=path.name,
                    is_scanned=is_scanned
                )
        except Exception as e:
            logger.error(f"pdfplumber also failed on {path.name}: {e}")
            raise RuntimeError(f"Failed to extract text from PDF '{path.name}': {e}")

    @classmethod
    def _load_pdf_stream(cls, stream: io.BytesIO, filename: str) -> LoadedDocument:
        """Extract text from in-memory PDF bytes with layout-aware spatial sorting."""
        try:
            import fitz
            doc = fitz.open(stream=stream.getvalue(), filetype="pdf")
            pages_text = [cls._extract_page_blocks_layout_aware(page) for page in doc]
            page_count = len(doc)
            doc.close()
            full_text = "\n\n".join(pages_text)
            total_chars = sum(len(p.strip()) for p in pages_text)
            is_scanned = (total_chars < 50 and page_count >= 1)

            metadata = {"engine": "pymupdf_spatial"}
            if is_scanned:
                metadata["is_scanned"] = True
                metadata["scanned_warning"] = "Low text density: document may be a scanned image or photograph."

            return LoadedDocument(
                raw_text=full_text,
                file_type="pdf",
                page_count=page_count,
                pages=pages_text,
                metadata=metadata,
                filename=filename,
                is_scanned=is_scanned
            )
        except Exception as e:
            logger.warning(f"PyMuPDF stream extraction failed: {e}. Trying pdfplumber.")

        try:
            import pdfplumber
            stream.seek(0)
            with pdfplumber.open(stream) as pdf:
                pages_text = [page.extract_text() or "" for page in pdf.pages]
                full_text = "\n\n".join(pages_text)
                total_chars = sum(len(p.strip()) for p in pages_text)
                is_scanned = (total_chars < 50 and len(pdf.pages) >= 1)

                metadata = {"engine": "pdfplumber"}
                if is_scanned:
                    metadata["is_scanned"] = True
                    metadata["scanned_warning"] = "Low text density: document may be a scanned image."

                return LoadedDocument(
                    raw_text=full_text,
                    file_type="pdf",
                    page_count=len(pdf.pages),
                    pages=pages_text,
                    metadata=metadata,
                    filename=filename,
                    is_scanned=is_scanned
                )
        except Exception as e:
            raise RuntimeError(f"Failed to extract text from PDF stream: {e}")

    # --- DOCX Processing ---

    @classmethod
    def _load_docx(cls, path: Path) -> LoadedDocument:
        """Extract text and table content from Microsoft Word document."""
        import docx
        doc = docx.Document(str(path))
        return cls._extract_docx_paragraphs_and_tables(doc, filename=path.name)

    @classmethod
    def _load_docx_stream(cls, stream: io.BytesIO, filename: str) -> LoadedDocument:
        """Extract text and tables from in-memory DOCX stream."""
        import docx
        doc = docx.Document(stream)
        return cls._extract_docx_paragraphs_and_tables(doc, filename=filename)

    @classmethod
    def _extract_docx_paragraphs_and_tables(cls, doc, filename: str) -> LoadedDocument:
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text to avoid missing education or experience in tables
        table_rows = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_rows.append(" | ".join(row_cells))
                    
        full_text = "\n".join(paragraphs)
        if table_rows:
            full_text += "\n\n" + "\n".join(table_rows)

        return LoadedDocument(
            raw_text=full_text,
            file_type="docx",
            page_count=1,
            pages=[full_text],
            metadata={"paragraphs_count": len(paragraphs), "tables_count": len(doc.tables)},
            filename=filename,
            is_scanned=False
        )

    # --- TXT Processing ---

    @classmethod
    def _load_txt(cls, path: Path) -> LoadedDocument:
        """Extract text from plaintext file handling multiple encodings."""
        data = path.read_bytes()
        return cls._load_txt_bytes(data, filename=path.name)

    @classmethod
    def _load_txt_bytes(cls, data: bytes, filename: str) -> LoadedDocument:
        """Decode raw text bytes across UTF-8, Latin-1, and CP1252 encodings."""
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                text = data.decode(encoding)
                return LoadedDocument(
                    raw_text=text,
                    file_type="txt",
                    page_count=1,
                    pages=[text],
                    metadata={"encoding": encoding},
                    filename=filename,
                    is_scanned=False
                )
            except UnicodeDecodeError:
                continue

        # Fallback with replacement
        text = data.decode("utf-8", errors="replace")
        return LoadedDocument(
            raw_text=text,
            file_type="txt",
            page_count=1,
            pages=[text],
            metadata={"encoding": "utf-8-replace"},
            filename=filename,
            is_scanned=False
        )
